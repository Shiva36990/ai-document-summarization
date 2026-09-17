import re
import numpy as np
import faiss
import nltk

from pypdf import PdfReader
from docx import Document
from nltk.tokenize import sent_tokenize

from sentence_transformers import SentenceTransformer

from transformers import (
    AutoTokenizer,
    AutoModelForSeq2SeqLM
)

# -----------------------------------
# DOWNLOAD NLTK TOKENIZER
# -----------------------------------

nltk.download('punkt')
nltk.download('punkt_tab')

# -----------------------------------
# LOAD EMBEDDING MODEL
# -----------------------------------

print("Loading embedding model...")

embedding_model = SentenceTransformer(
    'sentence-transformers/all-MiniLM-L6-v2'
)

# -----------------------------------
# LOAD SUMMARIZATION MODEL
# -----------------------------------

print("Loading summarization model...")

tokenizer = AutoTokenizer.from_pretrained(
    "facebook/bart-large-cnn"
)

model = AutoModelForSeq2SeqLM.from_pretrained(
    "facebook/bart-large-cnn"
)

# -----------------------------------
# EXTRACT TEXT
# -----------------------------------

def extract_text(file_path):
    text = ""
    file_path_lower = file_path.lower()

    if file_path_lower.endswith(".txt"):
        with open(file_path, "r", encoding="utf-8") as file:
            text = file.read()

    elif file_path_lower.endswith(".pdf"):
        reader = PdfReader(file_path)

        for page in reader.pages:
            extracted = page.extract_text()

            if extracted:
                text += extracted + "\n"

    elif file_path_lower.endswith(".docx"):
        document = Document(file_path)

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"

        # Extract text from tables also
        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    row_text.append(cell.text)

                text += " | ".join(row_text) + "\n"

    else:
        raise ValueError(
            "Unsupported file format. Please use .txt, .pdf, or .docx"
        )

    return text

# -----------------------------------
# CLEAN TEXT
# -----------------------------------

def clean_text(text):
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'\[\d+\]', '', text)
    text = re.sub(r'[^a-zA-Z0-9.,!? ]', '', text)
    return text.strip()

# -----------------------------------
# CHUNK TEXT
# -----------------------------------

def chunk_text(text, chunk_size=300):

    sentences = sent_tokenize(text)

    chunks = []

    current_chunk = []

    current_words = 0

    for sentence in sentences:

        word_count = len(sentence.split())

        if current_words + word_count <= chunk_size:

            current_chunk.append(sentence)

            current_words += word_count

        else:

            chunks.append(" ".join(current_chunk))

            current_chunk = [sentence]

            current_words = word_count

    if current_chunk:
        chunks.append(" ".join(current_chunk))

    return chunks

# -----------------------------------
# GENERATE EMBEDDINGS
# -----------------------------------

def generate_embeddings(chunks):

    embeddings = embedding_model.encode(
        chunks,
        convert_to_numpy=True
    )

    return embeddings.astype('float32')

# -----------------------------------
# BUILD FAISS INDEX
# -----------------------------------

def build_faiss_index(embeddings):

    dimension = embeddings.shape[1]

    index = faiss.IndexFlatL2(dimension)

    index.add(embeddings)

    return index

# -----------------------------------
# RETRIEVE IMPORTANT CHUNKS
# -----------------------------------

def retrieve_key_chunks(
    chunks,
    embeddings,
    index,
    top_k=5
):

    centroid = np.mean(
        embeddings,
        axis=0
    ).reshape(1, -1)

    distances, indices = index.search(
        centroid.astype('float32'),
        top_k
    )

    important_chunks = [
        chunks[i]
        for i in indices[0]
    ]

    return important_chunks

# -----------------------------------
# SUMMARIZE TEXT
# -----------------------------------

def summarize_text(text):

    inputs = tokenizer(
        text,
        max_length=1024,
        return_tensors="pt",
        truncation=True
    )

    summary_ids = model.generate(
        inputs["input_ids"],
        max_length=180,
        min_length=60,
        num_beams=4,
        early_stopping=True
    )

    summary = tokenizer.decode(
        summary_ids[0],
        skip_special_tokens=True
    )

    return summary

# -----------------------------------
# COMPLETE PIPELINE
# -----------------------------------

def document_summary(file_path):

    print("\nExtracting text...")

    raw_text = extract_text(file_path)

    print("Cleaning text...")

    cleaned_text = clean_text(raw_text)

    print("Chunking text...")

    chunks = chunk_text(cleaned_text)

    print(f"Total chunks created: {len(chunks)}")

    print("Generating embeddings...")

    embeddings = generate_embeddings(chunks)

    print("Building FAISS index...")

    index = build_faiss_index(embeddings)

    print("Retrieving important chunks...")

    important_chunks = retrieve_key_chunks(
        chunks,
        embeddings,
        index
    )

    combined_text = " ".join(important_chunks)

    print("Generating summary...\n")

    final_summary = summarize_text(combined_text)

    return final_summary

# -----------------------------------
# RUN PROGRAM
# -----------------------------------

if __name__ == "__main__":

    file_path = "Title The Impact of Excessive Mobil.txt"

    summary = document_summary(file_path)

    print("\nFINAL SUMMARY:\n")

    print(summary)